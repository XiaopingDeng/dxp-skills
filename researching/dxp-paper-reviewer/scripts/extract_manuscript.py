#!/usr/bin/env python3
"""提取稿件全文为纯文本，并报告页数与模板线索。

用法:
    python extract_manuscript.py <manuscript.pdf|docx> [-o out.txt]

输出:
    PAGES / CHARS / TEMPLATE_HINT / TEXT_FILE

PDF 走 pymupdf → pdfplumber → pypdf 依次降级；DOCX 走 python-docx → pandoc → 裸 XML。
"""
import argparse
import os
import subprocess
import sys
import zipfile
import re


def from_pdf(path):
    try:
        import fitz  # pymupdf
        doc = fitz.open(path)
        parts = []
        for i, page in enumerate(doc):
            parts.append("\n===== PAGE %d =====\n" % (i + 1) + page.get_text())
        return "".join(parts), doc.page_count, dict(doc.metadata or {})
    except ImportError:
        pass

    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            parts = []
            for i, page in enumerate(pdf.pages):
                parts.append("\n===== PAGE %d =====\n" % (i + 1) + (page.extract_text() or ""))
            return "".join(parts), len(pdf.pages), {}
    except ImportError:
        pass

    from pypdf import PdfReader
    reader = PdfReader(path)
    parts = []
    for i, page in enumerate(reader.pages):
        parts.append("\n===== PAGE %d =====\n" % (i + 1) + (page.extract_text() or ""))
    return "".join(parts), len(reader.pages), dict(reader.metadata or {})


def from_docx(path):
    try:
        import docx  # python-docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts), None, {}
    except ImportError:
        pass

    try:
        out = subprocess.run(["pandoc", path, "-t", "plain"],
                             capture_output=True, text=True, encoding="utf-8")
        if out.returncode == 0 and out.stdout.strip():
            return out.stdout, None, {}
    except FileNotFoundError:
        pass

    # 裸 XML 兜底：抓 <w:t> 文本
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "replace")
    xml = xml.replace("</w:p>", "\n")
    text = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))
    return text, None, {}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("-o", "--out", default=None)
    args = ap.parse_args()

    ext = os.path.splitext(args.path)[1].lower()
    if ext == ".pdf":
        text, pages, meta = from_pdf(args.path)
    elif ext in (".docx", ".doc"):
        text, pages, meta = from_docx(args.path)
    else:
        raise SystemExit("unsupported extension: %s" % ext)

    out = args.out or os.path.splitext(args.path)[0] + "_text.txt"
    with open(out, "w", encoding="utf-8") as f:
        f.write(text)

    print("PAGES:", pages if pages is not None else "n/a")
    print("CHARS:", len(text))
    print("TEMPLATE_HINT:", meta.get("creator"), "/", meta.get("producer"))
    print("TEXT_FILE:", out)
    sys.stdout.flush()


if __name__ == "__main__":
    main()
